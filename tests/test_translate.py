"""Offline tests for the Thai sets: translating new questions, checking and building them, the translator comparison,
and the Thai-aware parts of the later stages."""
import copy
import json
import math
import sys
import types

import pytest
import yaml

from distill import compress, data, generate, hub, rate, report, selection, translate, verify
from helpers import CFG, completion, gen_row, mock_client

SOURCE = [{"id": f"p{i}", "input": f"Name 3 fruits high in vitamin C, take {'abcdefghijkl'[i]}.", "strata": f"s{i % 2}", "closed": False,
           "reference": f"Reference {i}.", "hint_ok": i % 2 == 0, "source_dataset": "no_robots", "gold": []}
          for i in range(12)]
THAI = "บอกชื่อผลไม้ 3 ชนิดที่มีวิตามินซีสูง"


@pytest.fixture
def source(monkeypatch):
    monkeypatch.setattr(data, "load_prompts", lambda cfg, name: SOURCE)
    monkeypatch.setattr(data, "load_complexity", lambda ds: {f"p{i}": 2 for i in range(4)})  # used in English


# ---------------------------------------------------------------- choosing and translating the questions

def test_source_rows_skip_used_ids_and_nest(source):
    rows, pool, used = translate.source_rows(CFG, "no_robots_th", limit=6)
    assert (pool, used) == (8, 4) and not {r["id"] for r in rows} & {"p0", "p1", "p2", "p3"}
    first, _, _ = translate.source_rows(CFG, "no_robots_th", limit=3)
    assert first == rows[:3]  # the comparison's prompts are the first of the full run's


def test_source_rows_need_the_english_ratings_and_a_large_enough_pool(source, monkeypatch):
    with pytest.raises(ValueError, match="only 8 unused"):
        translate.source_rows(CFG, "no_robots_th", limit=9)
    monkeypatch.setattr(data, "load_complexity", lambda ds: {})
    with pytest.raises(FileNotFoundError, match="cannot be excluded"):
        translate.source_rows(CFG, "no_robots_th", limit=3)


def test_translate_requests_records_and_resume(source, tmp_path):
    out = tmp_path / "t.jsonl"
    bodies = []
    translate.run(CFG, "no_robots_th", "teacher", limit=4, out=out, client=mock_client(lambda p, b: completion(THAI),
                                                                                     bodies))
    sampling = CFG["translate"]["sampling"]
    assert {k: bodies[0][k] for k in sampling} == sampling and bodies[0]["model"] == CFG["teacher"]["vllm"]["model"]
    assert bodies[0]["chat_template_kwargs"] == {"enable_thinking": False}
    assert "<text>\nName 3" in bodies[0]["messages"][1]["content"] and "Thai" in bodies[0]["messages"][0]["content"]
    recs = report.read_jsonl(out)
    assert len(recs) == 4 and len(bodies) == 4 + sum(r["hint_ok"] for r in recs)  # references only when hint_ok
    r = recs[0]
    assert r["input"] == THAI and r["input_en"].startswith("Name 3") and r["dataset"] == "no_robots_th"
    assert r["translator"]["role"] == "teacher" and r["rule_flags"] == []
    assert all((x["reference"] == THAI) == x["hint_ok"] for x in recs)
    bodies.clear()
    translate.run(CFG, "no_robots_th", "teacher", limit=6, out=out, reference=False,
                  client=mock_client(lambda p, b: completion(THAI), bodies))
    assert len(bodies) == 2  # resumed: only the two new prompts, and no references


def test_typhoon_gets_its_card_prompt_and_dry_run_calls_nothing(source, tmp_path, capsys):
    bodies = []
    translate.run(CFG, "no_robots_th", "typhoon", limit=1, out=tmp_path / "t.jsonl", reference=False,
                  client=mock_client(lambda p, b: completion(THAI), bodies))
    (msg,) = bodies[0]["messages"]
    assert msg["role"] == "user" and msg["content"].startswith("Translate the following English text into Thai")
    translate.run(CFG, "no_robots_th", "teacher", limit=2, dry_run=True, client=mock_client(lambda p, b: 1 / 0))
    assert "4 prompts used in English excluded, 8 left, 2 sampled" in capsys.readouterr().out


def rec(en="Name 3 fruits high in vitamin C, from https://a.com.", th=THAI + " จาก https://a.com", **kw):
    return {"input_en": en, "input": th, "finish_reason": "stop", **kw}


@pytest.mark.parametrize("overrides, flag", [
    ({"finish_reason": "length"}, "not_stop"),
    ({"th": "Name 3 fruits high in vitamin C, from https://a.com."}, "not_thai"),
    ({"th": "ขออภัย ฉันไม่สามารถแปลข้อความนี้ได้ https://a.com 3"}, "refusal"),
    ({"th": "ผล 3 https://a.com", "en": "Name 3 fruits high in vitamin C. " * 10 + "https://a.com"}, "length_ratio"),
    ({"th": THAI.replace("3", "4") + " จาก https://a.com"}, "numbers_changed"),
    ({"en": "Run it:\n```print(1)```", "th": "รันโค้ดนี้:\n```พิมพ์(1)```"}, "code_changed"),
    ({"th": THAI + " จากเว็บไซต์"}, "urls_changed"),
    ({"en": "Which is true?\nA. One\nB. Two", "th": "ข้อใดเป็นจริง\nก. หนึ่ง\nข. สอง"}, "options_missing"),
    ({"en": "Translate into French: good morning", "th": "แปลเป็นภาษาฝรั่งเศส: good morning"}, "english_task_pattern"),
])
def test_rule_flags(overrides, flag):
    assert translate.rule_flags(rec(), CFG) == []
    kw = {k: v for k, v in overrides.items() if k not in ("en", "th")}
    r = rec(**{k: overrides[k] for k in ("en", "th") if k in overrides}, **kw)
    assert flag in translate.rule_flags(r, CFG)


def test_numbers_read_thai_digits_and_thousands_separators():
    assert not translate.numbers("Pay 1,000 baht in 2024") - translate.numbers("จ่าย ๑๐๐๐ บาท ในปี 2024")


# ---------------------------------------------------------------- the judge's check and embedding similarity

def digit(token, probs):
    return {"token": token, "top_logprobs": [{"token": t, "logprob": math.log(p)} for t, p in probs.items()]}


def test_check_parses_meaning_last_and_resumes(tmp_path):
    src = tmp_path / "t.jsonl"
    report.write_jsonl(src, [{"id": "a", "input_en": "Hi?", "input": "สวัสดี?"},
                             {"id": "b", "input_en": "Bye?", "input": ""}])  # an empty translation is not checked
    content = '{"english_only": false, "critique": "Fine.", "fluency": 4, "meaning": 4}'
    logprobs = {"content": [digit("4", {"4": 0.9, "9": 0.1}), digit("4", {"4": 0.6, "5": 0.4})]}
    bodies = []
    translate.check(CFG, src, client=mock_client(lambda p, b: completion(content, logprobs=logprobs), bodies))
    (out,) = report.read_jsonl(tmp_path / "t.check.jsonl")
    assert out["meaning"] == 4 and out["meaning_expected"] == pytest.approx(4.4) and out["english_only"] is False
    fmt = bodies[0]["response_format"]["json_schema"]
    assert fmt["name"] == "translation_check" and list(fmt["schema"]["properties"])[-1] == "meaning"
    assert "[English]\nHi?" in bodies[0]["messages"][1]["content"] and bodies[0]["temperature"] == 0
    bodies.clear()
    translate.check(CFG, src, client=mock_client(lambda p, b: completion(content), bodies))
    assert bodies == []


def test_check_rejects_a_reply_without_scores():
    with pytest.raises(ValueError, match="no valid scores"):
        translate.parse_check(completion('{"english_only": true, "critique": "x", "meaning": 7}'))


def test_similarity_with_a_fake_encoder(tmp_path):
    src = tmp_path / "t.jsonl"
    report.write_jsonl(src, [{"id": "a", "input_en": "same", "input": "same"},
                             {"id": "b", "input_en": "same", "input": "other"}])
    vec = {"same": [1.0, 0.0], "other": [0.0, 2.0]}
    translate.similarity(CFG, src, encode=lambda texts: [vec[t] for t in texts])
    assert {r["id"]: r["similarity"] for r in report.read_jsonl(tmp_path / "t.similarity.jsonl")} == {"a": 1.0, "b": 0.0}


# ---------------------------------------------------------------- build

def test_build_drops_failed_translations_and_rebuilds_option_golds(tmp_path, monkeypatch):
    monkeypatch.setattr(data, "PREPARED_DIR", tmp_path / "prepared")
    cfg = copy.deepcopy(CFG)
    cfg["translate"].update(min_meaning=4.0, min_similarity=0.7)
    th = "ข้อใดถูก\n\nA. หนึ่ง\nB. สอง\n\nตอบด้วยตัวอักษร"
    base = {"source_dataset": "logiqa", "input_en": "Which?\n\nA. One\nB. Two", "input": th, "strata": "x",
            "closed": True, "gold": ["B", "B. Two", "Two"], "hint_ok": True, "reference": None, "rule_flags": [],
            "translator": {"model": "qwen"}}
    path = tmp_path / "logiqa_th.jsonl"
    report.write_jsonl(path, [{**base, "id": i} for i in ("ok", "flagged", "english", "weak", "far", "unchecked")])
    path_rows = report.read_jsonl(path)
    path_rows[1]["rule_flags"] = ["not_thai"]
    report.write_jsonl(path, path_rows)
    ok = {"english_only": False, "fluency": 5, "meaning": 5, "meaning_expected": 4.9}
    report.write_jsonl(tmp_path / "logiqa_th.check.jsonl", [
        {"id": "ok", **ok}, {"id": "flagged", **ok}, {"id": "english", **ok, "english_only": True},
        {"id": "weak", **ok, "meaning_expected": 3.2}, {"id": "far", **ok}])
    report.write_jsonl(tmp_path / "logiqa_th.similarity.jsonl",
                       [{"id": i, "similarity": 0.5 if i == "far" else 0.9} for i in ("ok", "english", "weak", "far")])
    stats = translate.build(cfg, "logiqa_th", path)
    assert stats == {"translated": 6, "dropped: rule: not_thai": 1, "dropped: english_only": 1, "dropped: meaning": 1,
                     "dropped: similarity": 1, "dropped: unchecked": 1, "kept": 1}
    (row,) = report.read_jsonl(tmp_path / "prepared" / "logiqa_th.jsonl")
    assert row["gold"] == ["B", "B. สอง", "สอง"] and row["source_dataset"] == "logiqa_th" and row["lang"] == "th"
    assert row["closed"] and row["input_en"].startswith("Which?") and not row["hint_ok"]
    assert row["translation"] == {"model": "qwen", "meaning": 4.9, "fluency": 5, "similarity": 0.9}
    assert json.loads((tmp_path / "prepared" / "stats.json").read_text(encoding="utf-8"))["logiqa_th"]["kept"] == 1


def test_translated_sets_are_skipped_by_prepare_and_keep_exact_golds(tmp_path, monkeypatch):
    assert data.exact_gold(CFG) == {"logiqa", "logiqa_th"}
    monkeypatch.chdir(tmp_path)
    cfg = {"datasets": {k: CFG["datasets"][k] for k in ("logiqa", "logiqa_th")}, "verify": CFG["verify"]}
    row = {"text": json.dumps({"text": "P.", "question": "Q?", "options": ["a", "b"], "answer": 0,
                               "type": {"Categorical Reasoning": True}})}
    monkeypatch.setitem(sys.modules, "datasets", types.SimpleNamespace(load_dataset=lambda hf_id, split: [row]))
    data.prepare(cfg)
    assert sorted(p.name for p in data.PREPARED_DIR.iterdir()) == ["logiqa.jsonl", "stats.json"]
    with pytest.raises(ValueError, match="build-translated"):
        data.prepare(cfg, only="logiqa_th")


# ---------------------------------------------------------------- translator comparison

def comparison_files(tmp_path):
    """Two prompts, each translated by three translators, with checks and similarities."""
    paths = []
    for role, quality in (("teacher", 5), ("judge", 4), ("typhoon", 2)):
        path = tmp_path / f"no_robots_th.{role}.jsonl"
        report.write_jsonl(path, [{"id": f"q{i}", "dataset": "no_robots_th", "source_dataset": "no_robots",
                                   "input_en": f"English prompt {i}", "input": f"คำแปลที่ {i} แบบ {'กขค'[len(paths)]}",
                                   "translator": {"role": role, "model": f"model-{role}"},
                                   "rule_flags": ["not_thai"] if role == "typhoon" else []} for i in range(2)])
        report.write_jsonl(f"{str(path)[:-6]}.check.jsonl", [
            {"id": f"q{i}", "english_only": i == 1, "fluency": quality, "meaning": quality,
             "meaning_expected": quality - 0.1} for i in range(2)])
        report.write_jsonl(f"{str(path)[:-6]}.similarity.jsonl",
                           [{"id": f"q{i}", "similarity": 0.5 + quality / 10} for i in range(2)])
        paths.append(path)
    return paths


def test_compare_sheet_groups_prompts_and_hides_the_translators(tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "compare.xlsx"
    key = translate.compare_sheet(comparison_files(tmp_path), out)
    rows = list(load_workbook(out)["grades"].iter_rows(values_only=True))
    assert list(rows[0]) == translate.HEADER and len(rows) == 7
    english = [r[2] for r in rows[1:]]
    assert english[0] == english[1] == english[2] and english[3] == english[4] == english[5] != english[0]
    assert not any("model-" in str(v) or "teacher" in str(v) for r in rows for v in r)  # blind
    assert set(json.loads((tmp_path / "compare_key.json").read_text(encoding="utf-8"))) == {str(n) for n in key}
    assert {k["translator"] for k in key.values()} == {"teacher", "judge", "typhoon"}


def test_compare_report_scores_translators_and_signals(tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "compare.xlsx"
    key = translate.compare_sheet(comparison_files(tmp_path), out)
    wb = load_workbook(out)
    ws = wb["grades"]
    grade = {"teacher": 5, "judge": 4, "typhoon": 2}
    for row in ws.iter_rows(min_row=2):
        g = grade[key[row[0].value]["translator"]]
        row[4].value, row[5].value = g, g
    ws.cell(row=2, column=7, value="odd | phrase")
    wb.save(out)
    translate.compare_report(out, tmp_path / "report.md")
    text = (tmp_path / "report.md").read_text(encoding="utf-8")
    assert "6 of 6 translations graded" in text
    assert "| teacher | model-teacher | 2 | 5.00 | 5.00 | 0% | 0% | 50% | -0.10 |" in text
    assert "| typhoon | model-typhoon | 2 | 2.00 | 2.00 | 100% | 100% | 50% | -0.10 |" in text
    assert "| judge meaning (expected) | meaning | 6 | 1.00 | 1.00 |" in text
    assert "| embedding similarity | 0.75 | 100% | 100% |" in text  # the bad ones (0.7) out, the good (0.9, 1.0) in
    assert "| judge meaning (expected) | 4.0 | 100% | 50% |" in text  # the judge's own 3.9 falls below
    assert "| not_thai | 2 | 2.00 |" in text and "odd / phrase" in text and "English prompt 1" in text


def test_spearman_handles_ties():
    assert translate.ranks([3, 1, 3]) == [2.5, 1.0, 2.5]
    assert translate.spearman([1, 2, 3], [10, 20, 30]) == pytest.approx(1.0)
    assert translate.spearman([1, 1, 1], [1, 2, 3]) is None


# ---------------------------------------------------------------- Thai in the later stages

def test_rate_uses_the_english_original(tmp_path, monkeypatch):
    monkeypatch.setattr(data, "load_prompts", lambda cfg, name: [
        {"id": "a", "input": THAI, "input_en": "Name 3 fruits.", "strata": None, "closed": False}])
    bodies = []
    rate.run(CFG, "no_robots_th", out=tmp_path / "r.jsonl",
             client=mock_client(lambda p, b: completion("ReasoningRequired: 1"), bodies))
    assert "Name 3 fruits." in bodies[0]["messages"][-1]["content"] and THAI not in bodies[0]["messages"][-1]["content"]


def test_a_dataset_can_pick_its_own_prompt_versions(tmp_path, monkeypatch):
    cfg = copy.deepcopy(CFG)
    cfg["datasets"]["no_robots_th"]["prompt_versions"] = {"open": "th1", "compress": "th1"}
    monkeypatch.setattr(data, "load_prompts", lambda cfg, name: [
        {**SOURCE[0], "input": THAI, "input_en": SOURCE[0]["input"], "lang": "th", "translation": {"model": "q"}}])
    bodies = []
    client = mock_client(lambda p, b: completion("คำตอบ", reasoning="Because.", reasoning_tokens=5), bodies)
    generate.run(cfg, "no_robots_th", n=1, out=tmp_path / "g.jsonl", client=client)
    generate.run(cfg, "no_robots", n=1, out=tmp_path / "e.jsonl", client=client)
    assert "ตอบเป็นภาษาไทย" in bodies[0]["messages"][0]["content"]
    assert bodies[1]["messages"][0]["content"].startswith("You are a helpful assistant")
    sample = report.read_jsonl(tmp_path / "g.jsonl")[0]
    assert sample["prompt_version"] == "th1" and sample["lang"] == "th" and sample["input_en"].startswith("Name 3")
    src = tmp_path / "raw.jsonl"
    report.write_jsonl(src, [gen_row(0, source_dataset="no_robots_th"), gen_row(1)])
    bodies.clear()
    compress.run(cfg, src, tmp_path / "b.jsonl", client=mock_client(lambda p, b: completion("1. ขั้นตอน"), bodies))
    users = sorted(b["messages"][1]["content"][:6] for b in bodies)
    assert users == ["Questi", "คำถาม:"]
    hashes = {r["source_dataset"]: r["compress_prompt_hash"] for r in report.read_jsonl(tmp_path / "b.jsonl")}
    assert hashes["no_robots_th"] != hashes["no_robots"]


def test_thai_final_answers_and_words():
    assert verify.is_correct("เหตุผล...\nคำตอบสุดท้าย: B", ["B"])
    assert verify.is_correct("Final answer: ไม่ใช่", ["ไม่ใช่"]) and not verify.is_correct("Final answer: ใช่", ["ไม่ใช่"])
    assert verify.words("ฉันชอบกินข้าวผัด มากที่สุด") == ["ฉัน", "ชอบ", "กิน", "ข้าวผัด", "มาก", "ที่สุด"]
    assert verify.words("two words") == ["two", "words"] and verify.words(None) == []


def test_select_tells_thai_rationales_apart():
    a = selection.trigrams("1. แมวทุกตัวเป็นสัตว์เลี้ยงลูกด้วยนม\n\n2. ดังนั้นคำตอบคือ B")
    b = selection.trigrams("1. นกบางชนิดบินไม่ได้ เช่น นกเพนกวิน\n\n2. ดังนั้นคำตอบคือ C")
    assert len(a) > 5 and selection.jaccard(a, b) < 0.7


def test_filter_flags_an_english_rationale_for_a_thai_question():
    fcfg = CFG["filter"]
    row = gen_row(0, input="ช่วยอธิบายว่าทำไมท้องฟ้าถึงมีสีฟ้าในตอนกลางวัน และเป็นสีส้มในตอนเย็น",
                  answer="ท้องฟ้ามีสีฟ้าเพราะแสงสีน้ำเงินกระเจิงในบรรยากาศได้มากกว่าแสงสีอื่น ส่วนตอนเย็นแสงเดินทางไกลขึ้น",
                  rationale="1. Sunlight scatters off air molecules, and blue light scatters the most.\n\n"
                            "2. At sunset light travels through more air, so red and orange remain.",
                  cot_method="compressed")
    assert "language_mismatch" in report.rejection_reasons(row, fcfg)
    thai_steps = "1. แสงอาทิตย์กระเจิงเมื่อกระทบโมเลกุลอากาศ โดยแสงสีน้ำเงินกระเจิงมากที่สุด\n\n2. ตอนเย็นแสงผ่านอากาศมากขึ้น"
    assert report.rejection_reasons({**row, "rationale": thai_steps}, fcfg) == []


def test_final_rows_carry_the_translation():
    r = gen_row(0, source_dataset="logiqa_th", lang="th", input_en="Which?", translation={"model": "qwen"})
    final = selection.final_row(r, 2, {"method": "rule+judge"}, len)
    assert final["lang"] == "th" and final["input_en"] == "Which?" and final["translation"] == {"model": "qwen"}
    assert "lang" not in selection.final_row(gen_row(1), 2, {}, len)


def test_card_keeps_english_configs_and_names_thai_subsets(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    row = [{"id": "a", "complexity": 1, "input": "q", "rationale": "1. r", "answer": "a"}]
    for f in ("data/full/no_robots_open_B.final.jsonl", "data/full_th/no_robots_th_open_B.final.jsonl",
              "data/full_th/logiqa_th_closed_B.final.jsonl", "data/full_th/logiqa_th_closed_B_raw.jsonl"):
        (tmp_path / f).parent.mkdir(parents=True, exist_ok=True)
        report.write_jsonl(tmp_path / f, row)
    (tmp_path / "data/prepared").mkdir()
    text = hub.card("data/full_th", "org/data", runs=["data/full", "data/full_th"])
    meta = yaml.safe_load(text.split("---")[1])
    subsets = {c["config_name"]: [(f["split"], f["path"]) for f in c["data_files"]] for c in meta["configs"]}
    assert list(subsets) == ["final", "final_th", "final_th_closed"]  # raw files get no subset
    assert subsets["final"] == [("no_robots", "data/full/no_robots_open_B.final.jsonl")]
    assert subsets["final_th"] == [("no_robots", "data/full_th/no_robots_th_open_B.final.jsonl")]
    assert subsets["final_th_closed"] == [("logiqa", "data/full_th/logiqa_th_closed_B.final.jsonl")]
    assert meta["language"] == ["en", "th"] and "| logiqa (Thai, closed) | 1 | — | 1 |" in text
    assert "### no_robots (Thai): complexity 1" in text
